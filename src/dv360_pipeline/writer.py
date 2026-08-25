"""
Layer 5: fills the campaign burst Excel template from CampaignBurstData.

Formula-vs-static approach (confirmed with the user):
  - Row-level raw metrics (Impressions, TrueView, Clicks, Spends, video
    quartiles) are static values written by Python - they're source data.
  - Per-row View Rate / CTR are Excel formulas referencing that row's own
    Impressions/TrueView/Clicks cells (single-cell refs, safe even after
    row insertion).
  - Section Total rows are Excel SUM() formulas, rewritten with the correct
    row range once the section's actual size is known.
  - Total-row View Rate / CTR are formulas referencing the Total row's own
    SUM cells.
  - Pace / Ideal (row 7/9) are left untouched - they're already formulas in
    the template referencing single cells (Budget, Spend, flight/report
    dates) that we fill directly, so they stay valid regardless of how the
    breakdown sections below grow or shrink.

Breakdown sections (Creative, Targeting, Device, Gender, Age) are dynamically
resized to fit the actual number of distinct values found in the data,
inserting or deleting rows as needed and shifting everything below down/up.
The Date section is pre-formatted in the template for up to 56 rows with a
fixed-range Total formula; normal reporting windows fit inside that without
any resizing, so we only insert extra rows if the date range exceeds it.
"""

import datetime as _dt
import re
from copy import copy

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from dv360_pipeline.models import CampaignBurstData


def _to_date(value):
    """Normalize a pandas Timestamp / datetime to a plain date for Excel;
    pass other values through unchanged."""
    if value is None:
        return value
    if hasattr(value, "to_pydatetime"):
        value = value.to_pydatetime()
    if isinstance(value, _dt.datetime):
        return value.date()
    return value

IO_NAME_SHEET = "IO_name"
DATA_TEMPLATE_SHEET = "Data Template"

# Excel worksheet-name constraints.
EXCEL_SHEET_MAXLEN = 31
_EXCEL_ILLEGAL_CHARS = re.compile(r"[:\\/?*\[\]]")


def _io_sheet_label(io_name: str) -> str:
    """Sheet label for an IO = its io_name's last two '-'-delimited segments
    rejoined with '-' (e.g. '...-Video Reach Campaign-Burst 5' -> 'Video Reach
    Campaign-Burst 5'). Confirmed against the client's naming sheet."""
    parts = str(io_name).split("-")
    return "-".join(parts[-2:]).strip()


def _truncate_middle(text: str, limit: int) -> str:
    """Shorten to `limit` chars keeping both ends (so trailing distinguishers
    like 'Burst 7' survive), joining the halves with a single ellipsis."""
    if len(text) <= limit:
        return text
    if limit <= 1:
        return text[:limit]
    keep = limit - 1  # room for the ellipsis
    head = (keep + 1) // 2
    tail = keep - head
    return text[:head] + "…" + (text[-tail:] if tail else "")


def _excel_safe_sheet_name(label: str, used: set) -> str:
    """Make `label` a valid, unique Excel sheet name: strip illegal chars,
    fit within 31 chars, and de-duplicate case-insensitively with a numeric
    suffix (Excel treats sheet names case-insensitively for uniqueness)."""
    name = _EXCEL_ILLEGAL_CHARS.sub(" ", str(label)).strip() or "IO"
    name = _truncate_middle(name, EXCEL_SHEET_MAXLEN)
    if name.casefold() not in used:
        used.add(name.casefold())
        return name
    i = 2
    while True:
        suffix = f" ({i})"
        candidate = _truncate_middle(name, EXCEL_SHEET_MAXLEN - len(suffix)) + suffix
        if candidate.casefold() not in used:
            used.add(candidate.casefold())
            return candidate
        i += 1

CAMPAIGN_NAME_CELL = "C6"
FLIGHT_START_CELL = "C7"
FLIGHT_END_CELL = "D7"
CURRENCY_CELL = "C8"
BUDGET_CELL = "C9"
SPEND_CELL = "C10"
GUARANTEED_RATE_CELL = "C11"
KPI_CELL = "C12"
KPI_LABEL_CELL = "B12"
KPI_INVENTORY_LABEL_CELL = "B13"
KPI_INVENTORY_VALUE_CELL = "C13"  # template already merges C13:D13 with #,##0
KPI_MERGE = "B12:B13"
REPORT_START_CELL = "C14"
REPORT_END_CELL = "D14"

# All dates in the report render day/month/year.
DATE_FORMAT = "DD/MM/YYYY"

# name_col: 1-indexed column for the row label (B=2)
NAME_COL = 2
IMPRESSIONS_COL = 3   # C
TRUEVIEW_COL = 4      # D
SPEND_COL = 5          # E
VIEW_RATE_COL = 6     # F
CLICKS_COL = 7         # G
CTR_COL = 8             # H
Q25_COL = 9             # I
Q50_COL = 10            # J
Q75_COL = 11            # K
Q100_COL = 12           # L

BASE_METRIC_COLS = [IMPRESSIONS_COL, TRUEVIEW_COL, SPEND_COL, CLICKS_COL]
QUARTILE_COLS = [Q25_COL, Q50_COL, Q75_COL, Q100_COL]


def _copy_row_style(ws: Worksheet, src_row: int, dst_row: int, min_col: int, max_col: int) -> None:
    for col in range(min_col, max_col + 1):
        src_cell = ws.cell(row=src_row, column=col)
        dst_cell = ws.cell(row=dst_row, column=col)
        dst_cell.number_format = src_cell.number_format
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.alignment = copy(src_cell.alignment)


def _resize_section(
    ws: Worksheet,
    header_row: int,
    total_row_template: int,
    actual_n: int,
    fixed_capacity: bool,
    max_col: int,
) -> tuple[int, int]:
    """
    Returns (data_start_row, new_total_row). Inserts or deletes rows so the
    section has exactly `actual_n` data rows (unless fixed_capacity and
    actual_n already fits within the template's pre-built capacity).
    """
    data_start = header_row + 1
    template_n = total_row_template - header_row - 1

    if fixed_capacity and actual_n <= template_n:
        return data_start, total_row_template

    if actual_n > template_n:
        n_insert = actual_n - template_n
        ws.insert_rows(total_row_template, amount=n_insert)
        last_template_data_row = total_row_template - 1
        for i in range(n_insert):
            _copy_row_style(ws, last_template_data_row, last_template_data_row + 1 + i, NAME_COL, max_col)
    elif actual_n < template_n:
        n_delete = template_n - actual_n
        delete_from = data_start + actual_n
        ws.delete_rows(delete_from, amount=n_delete)

    new_total_row = header_row + actual_n + 1
    return data_start, new_total_row


def _write_section(
    ws: Worksheet,
    header_row: int,
    total_row_template: int,
    df: pd.DataFrame,
    name_column: str,
    has_quartiles: bool,
    fixed_capacity: bool = False,
) -> int:
    """Writes one breakdown section. Returns the row-count delta caused (for cascading offsets)."""
    max_col = Q100_COL if has_quartiles else CTR_COL
    actual_n = len(df)

    data_start, total_row = _resize_section(ws, header_row, total_row_template, actual_n, fixed_capacity, max_col)

    for i, row in enumerate(df.itertuples(index=False)):
        r = data_start + i
        name_value = _to_date(getattr(row, name_column))
        name_cell = ws.cell(row=r, column=NAME_COL, value=name_value)
        if isinstance(name_value, _dt.date):  # Date breakdown -> DD/MM/YYYY
            name_cell.number_format = DATE_FORMAT
        ws.cell(row=r, column=IMPRESSIONS_COL, value=int(row.impressions))
        ws.cell(row=r, column=TRUEVIEW_COL, value=int(row.trueview_views))
        ws.cell(row=r, column=SPEND_COL, value=round(float(row.spend), 2))
        ws.cell(row=r, column=CLICKS_COL, value=int(row.clicks))
        ws.cell(row=r, column=VIEW_RATE_COL).value = f"=D{r}/C{r}"
        ws.cell(row=r, column=CTR_COL).value = f"=G{r}/C{r}"
        if has_quartiles:
            ws.cell(row=r, column=Q25_COL, value=int(row.video_q25))
            ws.cell(row=r, column=Q50_COL, value=int(row.video_q50))
            ws.cell(row=r, column=Q75_COL, value=int(row.video_q75))
            ws.cell(row=r, column=Q100_COL, value=int(row.video_q100))

    # Always rewrite the Total row's formulas against the *current* row
    # positions. Even when a section isn't resized (e.g. the Date section
    # fits within its pre-built capacity), everything below any earlier
    # section that *was* resized has physically shifted - but openpyxl does
    # not adjust formula text for cells elsewhere in the sheet, so a
    # formula left untouched would keep pointing at stale template rows.
    last_data_row = data_start + actual_n - 1 if actual_n > 0 else data_start
    ws.cell(row=total_row, column=NAME_COL, value="Total")
    for col in [IMPRESSIONS_COL, TRUEVIEW_COL, SPEND_COL, CLICKS_COL] + (QUARTILE_COLS if has_quartiles else []):
        letter = get_column_letter(col)
        ws.cell(row=total_row, column=col).value = f"=SUM({letter}{data_start}:{letter}{last_data_row})"
    ws.cell(row=total_row, column=VIEW_RATE_COL).value = f"=D{total_row}/C{total_row}"
    ws.cell(row=total_row, column=CTR_COL).value = f"=G{total_row}/C{total_row}"

    return total_row - total_row_template


def _copy_sheet_view(src_ws: Worksheet, dst_ws: Worksheet) -> None:
    """openpyxl's copy_worksheet drops some sheet-view display settings — most
    visibly `showGridLines` (the template hides gridlines; copies show them,
    which makes the header box look cluttered) and the zoom level. Carry them
    over so every copied report sheet renders like the first one."""
    dst_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
    dst_ws.sheet_view.zoomScale = src_ws.sheet_view.zoomScale


def _split_kpi_inventory_row(ws: Worksheet, kpi_inventory) -> None:
    """Turn the merged KPI label block (B12:B13) into two rows: 'KPI:' (row 12)
    and a new 'KPI Inventory:' (row 13). The value cell C13:D13 already exists
    in the template (merged, #,##0), so only the B13 label + value are added."""
    if KPI_MERGE in {str(m) for m in ws.merged_cells.ranges}:
        ws.unmerge_cells(KPI_MERGE)
    label = ws[KPI_INVENTORY_LABEL_CELL]
    src = ws[KPI_LABEL_CELL]
    label.value = "KPI Inventory:"
    label.font = copy(src.font)
    label.alignment = copy(src.alignment)
    # Unmerging clears B13's border, breaking the box's left edge — restore it
    # from the KPI label cell (which carries the left box border).
    label.border = copy(src.border)
    if kpi_inventory is not None:
        ws[KPI_INVENTORY_VALUE_CELL] = kpi_inventory


def _fill_io_name_sheet(ws: Worksheet, data: CampaignBurstData) -> None:
    """Fill one IO_name-style report sheet (header cells + the six breakdown
    sections) in place. Shared by the single- and multi-IO writers."""
    ws[CAMPAIGN_NAME_CELL] = data.meta.campaign_name
    ws[FLIGHT_START_CELL] = data.meta.flight_start
    ws[FLIGHT_END_CELL] = data.meta.flight_end
    ws[CURRENCY_CELL] = data.meta.currency
    ws[BUDGET_CELL] = data.meta.budget
    ws[SPEND_CELL] = round(data.spend, 2)
    ws[GUARANTEED_RATE_CELL] = data.meta.guaranteed_rate
    ws[KPI_CELL] = data.meta.kpi
    _split_kpi_inventory_row(ws, data.meta.kpi_inventory)

    # Reporting range: start from the requested window, but the end reflects
    # the last date that actually has data (the flight may outrun delivery).
    ws[REPORT_START_CELL] = data.report_start
    report_end = data.report_end
    if not data.date_df.empty and "date" in data.date_df.columns:
        report_end = _to_date(max(data.date_df["date"]))
    ws[REPORT_END_CELL] = report_end

    # All header dates render DD/MM/YYYY.
    for cell in (FLIGHT_START_CELL, FLIGHT_END_CELL, REPORT_START_CELL, REPORT_END_CELL):
        ws[cell].number_format = DATE_FORMAT

    offset = 0
    delta = _write_section(ws, 17 + offset, 19 + offset, data.creative_df, "creative_name", has_quartiles=True)
    offset += delta
    delta = _write_section(ws, 21 + offset, 26 + offset, data.targeting_df, "targeting", has_quartiles=False)
    offset += delta
    delta = _write_section(ws, 28 + offset, 33 + offset, data.device_df, "device_type", has_quartiles=False)
    offset += delta
    delta = _write_section(ws, 35 + offset, 38 + offset, data.gender_df, "gender", has_quartiles=False)
    offset += delta
    delta = _write_section(ws, 40 + offset, 43 + offset, data.age_df, "age", has_quartiles=False)
    offset += delta
    _write_section(ws, 45 + offset, 102 + offset, data.date_df, "date", has_quartiles=False, fixed_capacity=True)


def write_report(data: CampaignBurstData, template_path: str, output_path: str) -> None:
    """Single-IO report: fill the template's IO_name and Data Template sheets."""
    import openpyxl

    wb = openpyxl.load_workbook(template_path)
    _fill_io_name_sheet(wb[IO_NAME_SHEET], data)
    _write_data_template_sheet(wb, [data])
    wb.save(output_path)


def write_multi_io_report(datas: list[CampaignBurstData], template_path: str, output_path: str) -> None:
    """Multi-IO report: one IO_name-style sheet per IO (named by
    _io_sheet_label), plus a single combined Data Template with every IO's
    detail rows, and the static Sheet1 left untouched.

    Falls back to the single-IO writer for one IO. Report sheets are ordered
    first, then Data Template, then Sheet1.
    """
    import openpyxl

    if not datas:
        raise ValueError("datas must not be empty")
    if len(datas) == 1:
        return write_report(datas[0], template_path, output_path)

    wb = openpyxl.load_workbook(template_path)
    template_io = wb[IO_NAME_SHEET]

    # Copy the pristine template sheet once per IO *before* filling any of
    # them, so every copy starts from the unfilled template.
    io_sheets = [template_io] + [wb.copy_worksheet(template_io) for _ in range(len(datas) - 1)]

    # Seed with the retained sheets so an IO's computed name can't collide.
    used_names: set = {DATA_TEMPLATE_SHEET.casefold(), "sheet1"}
    for ws, data in zip(io_sheets, datas):
        ws.title = _excel_safe_sheet_name(_io_sheet_label(data.meta.io_name), used_names)
        _copy_sheet_view(template_io, ws)
        _fill_io_name_sheet(ws, data)

    _write_data_template_sheet(wb, datas)

    # Order: all IO report sheets, then Data Template, then Sheet1.
    tail = [wb[name] for name in (DATA_TEMPLATE_SHEET, "Sheet1") if name in wb.sheetnames]
    wb._sheets = io_sheets + tail

    wb.save(output_path)


def _write_data_template_sheet(wb, datas: list[CampaignBurstData]) -> None:
    """Fill the Data Template sheet, stacking every IO's detail rows below the
    header (row 2). One IO for the single-IO report, all IOs for a multi-IO one."""
    ws = wb[DATA_TEMPLATE_SHEET]
    row = 3  # header is row 2
    for data in datas:
        for r in data.data_template_df.itertuples(index=False):
            date_cell = ws.cell(row=row, column=1, value=_to_date(r.date))  # Date
            date_cell.number_format = DATE_FORMAT
            ws.cell(row=row, column=7, value=data.meta.campaign_name)    # Campaign Name
            ws.cell(row=row, column=9, value=r.creative_name)            # Creative
            ws.cell(row=row, column=10, value=r.targeting)               # Strategy
            ws.cell(row=row, column=15, value=round(float(r.spend), 2))  # Cost
            ws.cell(row=row, column=16, value=int(r.impressions))        # Impressions
            ws.cell(row=row, column=17, value=int(r.clicks))             # Clicks
            ws.cell(row=row, column=20, value=int(r.trueview_views))     # Video Views
            ws.cell(row=row, column=21, value=int(r.video_q25))          # 25% Completed View
            ws.cell(row=row, column=22, value=int(r.video_q50))          # 50% Completed View
            ws.cell(row=row, column=23, value=int(r.video_q75))          # 75% Completed View
            ws.cell(row=row, column=24, value=int(r.video_q100))         # 100% Completed View
            row += 1
