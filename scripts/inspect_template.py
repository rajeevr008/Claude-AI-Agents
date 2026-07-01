"""
Prints every non-empty cell (value + type, formula vs literal) and merged
cell ranges for each sheet in the campaign burst Excel template. Used during
Stage 1 to confirm the template's structure before writing the Layer 5 writer.

Usage:
    uv run scripts/inspect_template.py [path/to/template.xlsx]
"""

import sys

import openpyxl

DEFAULT_TEMPLATE = "templates/campaign_burst_template.xlsx"


def inspect(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    print("SHEETS:", wb.sheetnames)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n===== SHEET: {sheet_name} (dims: {ws.dimensions}) =====")
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
                tag = "FORMULA" if is_formula else cell.data_type
                print(f"  {cell.coordinate}: [{tag}] {cell.value!r}")

        if ws.merged_cells.ranges:
            print(f"\n  MERGED CELLS: {list(ws.merged_cells.ranges)}")


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_TEMPLATE
    inspect(path)
